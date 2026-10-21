import asyncio
import json
import random
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError


# ============================================================
# CONFIGURATION
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

CONTACTS_FILE = BASE_DIR / "contacts.xlsx"
SCREENSHOT_DIR = BASE_DIR / "screenshots"

JSON_REPORT = BASE_DIR / "whatsapp_report.json"
EXCEL_REPORT = BASE_DIR / "whatsapp_report.xlsx"

WHATSAPP_URL = "https://web.whatsapp.com"

SCREENSHOT_DIR.mkdir(exist_ok=True)


# ============================================================
# RANDOM DELAY
# ============================================================

async def random_delay():
    """Wait 2-5 seconds between actions."""
    delay = random.uniform(2, 5)
    print(f"Waiting {delay:.1f} seconds...")
    await asyncio.sleep(delay)


# ============================================================
# READ CONTACTS FROM EXCEL
# ============================================================

def read_contacts():

    if not CONTACTS_FILE.exists():
        raise FileNotFoundError(
            "contacts.xlsx was not found."
        )

    workbook = load_workbook(
        CONTACTS_FILE,
        data_only=True
    )

    sheet = workbook.active

    rows = list(
        sheet.iter_rows(values_only=True)
    )

    if not rows:
        raise ValueError(
            "contacts.xlsx is empty."
        )

    headers = [
        str(value).strip()
        if value is not None
        else ""
        for value in rows[0]
    ]

    required_columns = [
        "Name",
        "Phone",
        "Message"
    ]

    for column in required_columns:
        if column not in headers:
            raise ValueError(
                f"Missing column: {column}"
            )

    contacts = []

    for row in rows[1:]:

        values = dict(
            zip(headers, row)
        )

        name = str(
            values.get("Name") or ""
        ).strip()

        phone = str(
            values.get("Phone") or ""
        ).strip()

        message = str(
            values.get("Message") or ""
        ).strip()

        if not name or not phone:
            continue

        if not message:
            message = (
                "Hello {name}, "
                "this is a test message."
            )

        contacts.append({
            "name": name,
            "phone": phone,
            "message": message
        })

    return contacts


# ============================================================
# PERSONALIZE MESSAGE
# ============================================================

def personalize_message(template, name):

    return template.replace(
        "{name}",
        name
    )


# ============================================================
# CLEAN PHONE NUMBER
# ============================================================

def clean_phone_number(phone):

    return re.sub(
        r"[^\d]",
        "",
        phone
    )


# ============================================================
# SAVE JSON REPORT
# ============================================================

def save_json_report(results):

    with open(
        JSON_REPORT,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            results,
            file,
            indent=4,
            ensure_ascii=False
        )


# ============================================================
# SAVE EXCEL REPORT
# ============================================================

def save_excel_report(results):

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "WhatsApp Report"

    headers = [
        "Name",
        "Phone",
        "Message",
        "Status",
        "Error",
        "Last 3 Messages",
        "Timestamp",
        "Screenshot"
    ]

    sheet.append(headers)

    for result in results:

        last_messages = "\n".join(
            result.get(
                "last_3_messages",
                []
            )
        )

        sheet.append([
            result.get("name", ""),
            result.get("phone", ""),
            result.get("message", ""),
            result.get("status", ""),
            result.get("error", ""),
            last_messages,
            result.get("timestamp", ""),
            result.get("screenshot", "")
        ])

    # Make the report readable
    column_widths = {
        "A": 20,
        "B": 20,
        "C": 50,
        "D": 15,
        "E": 40,
        "F": 60,
        "G": 25,
        "H": 50
    }

    for column, width in column_widths.items():

        sheet.column_dimensions[
            column
        ].width = width

    workbook.save(
        EXCEL_REPORT
    )


# ============================================================
# WAIT FOR WHATSAPP LOGIN
# ============================================================

async def wait_for_whatsapp_login(page):

    print()
    print("=" * 60)
    print("WHATSAPP LOGIN")
    print("=" * 60)

    print()
    print("If QR code is displayed:")
    print()
    print("1. Open WhatsApp on your phone")
    print("2. Go to Settings")
    print("3. Select Linked Devices")
    print("4. Select Link a Device")
    print("5. Scan the QR code")
    print()

    try:

        # The search box appears after successful login.
        await page.wait_for_selector(
            'div[contenteditable="true"]',
            timeout=120000
        )

        print(
            "WhatsApp login detected."
        )

        return True

    except PlaywrightTimeoutError:

        print(
            "Login timeout."
        )

        return False


# ============================================================
# OPEN CONTACT
# ============================================================

async def open_contact(page, phone):

    clean_phone = clean_phone_number(
        phone
    )

    print(
        f"Opening phone number: {phone}"
    )

    url = (
        "https://web.whatsapp.com/send"
        f"?phone={clean_phone}"
    )

    await page.goto(
        url,
        wait_until="domcontentloaded"
    )

    await page.wait_for_timeout(
        5000
    )

    return clean_phone


# ============================================================
# CHECK CONTACT
# ============================================================

async def check_contact(page):

    possible_errors = [
        "Phone number shared via url is invalid",
        "not on WhatsApp",
        "phone number isn't on WhatsApp"
    ]

    for error_text in possible_errors:

        try:

            await page.wait_for_selector(
                f"text={error_text}",
                timeout=2500
            )

            return False

        except PlaywrightTimeoutError:
            pass

    return True


# ============================================================
# FIND MESSAGE BOX
# ============================================================

async def find_message_box(page):

    message_box = page.locator(
        'div[contenteditable="true"]'
    ).last

    await message_box.wait_for(
        state="visible",
        timeout=15000
    )

    return message_box


# ============================================================
# SEND MESSAGE
# ============================================================

async def send_message(
    page,
    contact
):

    name = contact["name"]
    phone = contact["phone"]

    message = personalize_message(
        contact["message"],
        name
    )

    timestamp = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    result = {
        "name": name,
        "phone": phone,
        "message": message,
        "status": "Failed",
        "error": "",
        "last_3_messages": [],
        "timestamp": timestamp,
        "screenshot": ""
    }

    print()
    print("=" * 60)
    print(f"CONTACT: {name}")
    print(f"PHONE:   {phone}")
    print("=" * 60)

    try:

        # ----------------------------------------------------
        # Open contact
        # ----------------------------------------------------

        clean_phone = await open_contact(
            page,
            phone
        )

        await random_delay()

        # ----------------------------------------------------
        # Check whether number exists
        # ----------------------------------------------------

        contact_exists = await check_contact(
            page
        )

        if not contact_exists:

            result["error"] = (
                "Contact not found or "
                "number is not registered on WhatsApp"
            )

            print(
                "Contact not found."
            )

            return result

        # ----------------------------------------------------
        # Find message box
        # ----------------------------------------------------

        print(
            "Waiting for message box..."
        )

        message_box = await find_message_box(
            page
        )

        await random_delay()

        # ----------------------------------------------------
        # Type message
        # ----------------------------------------------------

        print(
            f"Sending: {message}"
        )

        await message_box.click()

        await message_box.fill(
            message
        )

        await random_delay()

        # ----------------------------------------------------
        # Send
        # ----------------------------------------------------

        await message_box.press(
            "Enter"
        )

        print(
            "Message sent."
        )

        await page.wait_for_timeout(
            3000
        )

        # ----------------------------------------------------
        # Screenshot
        # ----------------------------------------------------

        safe_name = re.sub(
            r"[^a-zA-Z0-9_-]",
            "_",
            name
        )

        screenshot_file = (
            SCREENSHOT_DIR
            / f"{safe_name}_{clean_phone}.png"
        )

        await page.screenshot(
            path=str(screenshot_file),
            full_page=False
        )

        result["screenshot"] = (
            screenshot_file.name
        )

        print(
            f"Screenshot saved: "
            f"{screenshot_file.name}"
        )

        # ----------------------------------------------------
        # Extract last 3 messages
        # ----------------------------------------------------

        print(
            "Extracting last 3 messages..."
        )

        await random_delay()

        message_elements = page.locator(
            'div.message-in, div.message-out'
        )

        count = await message_elements.count()

        start = max(
            0,
            count - 3
        )

        last_messages = []

        for index in range(
            start,
            count
        ):

            try:

                element = (
                    message_elements.nth(
                        index
                    )
                )

                text_element = (
                    element.locator(
                        "span.selectable-text"
                    ).last
                )

                text = await text_element.inner_text(
                    timeout=3000
                )

                if text.strip():

                    last_messages.append(
                        text.strip()
                    )

            except Exception:

                continue

        result[
            "last_3_messages"
        ] = last_messages

        result["status"] = "Sent"

        print(
            "Last 3 messages extracted:"
        )

        for message_text in last_messages:
            print(
                f"  - {message_text}"
            )

    except Exception as error:

        result["error"] = str(
            error
        )

        print(
            f"Error processing {name}:"
        )

        print(error)

    return result


# ============================================================
# MAIN AUTOMATION
# ============================================================

async def main():

    print()
    print("=" * 60)
    print(
        "WHATSAPP MESSAGE SENDER + "
        "SMART DATA EXTRACTOR"
    )
    print("=" * 60)

    # --------------------------------------------------------
    # Read contacts
    # --------------------------------------------------------

    try:

        contacts = read_contacts()

    except Exception as error:

        print(
            f"Could not read contacts.xlsx: "
            f"{error}"
        )

        return

    print()
    print(
        f"Contacts loaded: "
        f"{len(contacts)}"
    )

    if not contacts:

        print(
            "No contacts to process."
        )

        return

    # --------------------------------------------------------
    # Start Playwright
    # --------------------------------------------------------

    async with async_playwright() as p:

        print()
        print(
            "Starting Chromium..."
        )

        browser = await p.chromium.launch(
            headless=False,
            slow_mo=100
        )

        context = await browser.new_context()

        page = await context.new_page()

        # ----------------------------------------------------
        # Open WhatsApp
        # ----------------------------------------------------

        print(
            "Opening WhatsApp Web..."
        )

        await page.goto(
            WHATSAPP_URL,
            wait_until="domcontentloaded"
        )

        # ----------------------------------------------------
        # Login
        # ----------------------------------------------------

        logged_in = (
            await wait_for_whatsapp_login(
                page
            )
        )

        if not logged_in:

            await browser.close()

            return

        # ----------------------------------------------------
        # Process each contact
        # ----------------------------------------------------

        results = []

        for contact in contacts:

            result = await send_message(
                page,
                contact
            )

            results.append(
                result
            )

            # Save after every contact
            save_json_report(
                results
            )

            save_excel_report(
                results
            )

            await random_delay()

        # ----------------------------------------------------
        # Final reports
        # ----------------------------------------------------

        save_json_report(
            results
        )

        save_excel_report(
            results
        )

        print()
        print("=" * 60)
        print("AUTOMATION COMPLETED")
        print("=" * 60)

        print()
        print(
            f"JSON report:"
        )
        print(JSON_REPORT)

        print()
        print(
            f"Excel report:"
        )
        print(EXCEL_REPORT)

        print()
        print(
            f"Screenshots:"
        )
        print(SCREENSHOT_DIR)

        print()

        # Keep browser open for 5 seconds
        await page.wait_for_timeout(
            5000
        )

        await browser.close()


# ============================================================
# START PROGRAM
# ============================================================

if __name__ == "__main__":

    asyncio.run(main())
